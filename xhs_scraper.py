#!/usr/bin/env python3
"""
小紅書童裝供應商爬蟲 - 帶暗語私訊功能版本
直接運行: python xhs_scraper.py

使用前請先：
1. pip install playwright requests openpyxl pandas
2. playwright install chromium
3. 準備好小紅書 Cookie
"""

import asyncio
import json
import re
import os
import csv
import random
import time
import logging
from pathlib import Path
from datetime import datetime
from dataclasses import dataclass, field, asdict
from typing import Optional, List, Dict
from enum import Enum
from urllib.parse import quote

# ========== 配置 ==========
COOKIE = ""  # 在這裡貼上你的 Cookie，或運行時輸入

KEYWORDS = [
    "童裝代理",
    "童裝批發招商",
    "婴儿连体衣批發",
    "宝宝衣服代理",
    "大牌童装批發",
    "婴儿服一件代發",
    "童装套装招商",
    "婴儿衣服货源",
]

MAX_PAGES = 2  # 每個關鍵字最大頁數
MAX_SUPPLIERS = 50  # 最大供應商數量
HEADLESS = False  # 設為 False 可以看到瀏覽器操作 (調試用)

# ========== 暗語私訊配置 ==========
ENABLE_MESSAGE = True  # 是否啟用私訊功能
MAX_MESSAGES_PER_RUN = 5  # 每次運行最多發送幾條私訊 (低頻率)
MESSAGE_DELAY_MIN = 300  # 發私訊最小間隔 (秒) = 5分鐘
MESSAGE_DELAY_MAX = 600  # 發私訊最大間隔 (秒) = 10分鐘

# 暗語訊息模板列表 (隨機選擇)
SECRET_MESSAGES = [
    # 類型1: 加微信暗語
    "你好～看到你家童裝很不錯，可以伽V聊聊合作嗎？",
    "親，想了解批發價，方便加嘉嗎？",
    "你好，對你們家童裝有興趣，能加薇聊嗎？",
    "這款真好看！想拿貨，可以夾V嗎？",
    "想做代理，方便加徽詳聊嗎？",
    
    # 類型2: 看主頁暗語
    "有興趣合作，煮頁有聯繫方式嗎？",
    "想了解更多，看主有詳細嗎？",
    "求合作～煮叶留了聯繫方式嗎？",
    
    # 類型3: 私信暗語
    "想批發，可以si我詳情嗎？",
    "對童裝很有興趣，絲我可以聊嗎？",
    "做代購的，想斯聊批發價～",
    
    # 類型4: 懂的來暗語
    "香港批發商，薰的來，想長期合作",
    "想做童裝批發，懂的懂，求私",
    "同行求拿貨渠道，薰的来聯繫",
    
    # 類型5: 一般詢問 + 暗語
    "你好！童裝質量看起來不錯，我是香港的，想了解代理方式，方便伽薇聊聊嗎？",
    "親愛的～我在做童裝批發，對你家款式很有興趣，可以看煮頁嗎？",
    "你好，我是做跨境電商的，想長期合作拿貨，si我可以詳聊嗎？",
]

# 數據文件路徑
DATA_FILE = "suppliers_data.csv"

# ========== 日誌 ==========
# 設為 DEBUG 可看更多細節
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(message)s')
logger = logging.getLogger()

# ========== 數據模型 ==========
class SupplierType(Enum):
    UNKNOWN = "未知"
    FACTORY = "工廠/廠家"
    BRAND = "品牌方"
    DISTRIBUTOR = "經銷商/代理"

class ContactType(Enum):
    WECHAT = "微信"
    PHONE = "電話"

@dataclass
class ContactInfo:
    type: ContactType
    value: str
    source: str = ""

@dataclass
class NoteStats:
    likes: int = 0
    comments: int = 0
    collects: int = 0
    
    @property
    def total(self) -> int:
        return self.likes + self.comments * 3 + self.collects * 2

@dataclass
class Note:
    note_id: str
    title: str
    content: str = ""
    author_id: str = ""
    author_name: str = ""
    stats: NoteStats = field(default_factory=NoteStats)
    has_supplier_kw: bool = False

@dataclass
class Author:
    user_id: str
    nickname: str
    bio: str = ""
    followers: int = 0
    notes_count: int = 0
    ip_location: str = ""
    
    @property
    def url(self) -> str:
        return f"https://www.xiaohongshu.com/user/profile/{self.user_id}"

@dataclass
class Supplier:
    user_id: str
    nickname: str
    bio: str = ""
    followers: int = 0
    notes_count: int = 0
    ip_location: str = ""
    supplier_type: str = "未知"
    wechat: str = ""
    score: float = 0.0
    profile_url: str = ""
    keywords: str = ""
    # 私訊相關
    message_sent: bool = False
    message_content: str = ""
    message_time: str = ""
    # 回覆相關
    reply_received: bool = False
    reply_content: str = ""
    reply_time: str = ""
    # 元數據
    created_at: str = ""
    updated_at: str = ""

# ========== 數據管理器 ==========
class DataManager:
    """管理供應商數據的持久化存儲"""
    
    def __init__(self, filepath: str = DATA_FILE):
        self.filepath = filepath
        self.suppliers: Dict[str, Supplier] = {}
        self.load()
    
    def load(self):
        """從CSV加載數據"""
        if not os.path.exists(self.filepath):
            return
        try:
            with open(self.filepath, 'r', encoding='utf-8') as f:
                reader = csv.DictReader(f)
                for row in reader:
                    s = Supplier(
                        user_id=row.get('user_id', ''),
                        nickname=row.get('nickname', ''),
                        bio=row.get('bio', ''),
                        followers=int(row.get('followers', 0) or 0),
                        notes_count=int(row.get('notes_count', 0) or 0),
                        ip_location=row.get('ip_location', ''),
                        supplier_type=row.get('supplier_type', '未知'),
                        wechat=row.get('wechat', ''),
                        score=float(row.get('score', 0) or 0),
                        profile_url=row.get('profile_url', ''),
                        keywords=row.get('keywords', ''),
                        message_sent=row.get('message_sent', '').lower() == 'true',
                        message_content=row.get('message_content', ''),
                        message_time=row.get('message_time', ''),
                        reply_received=row.get('reply_received', '').lower() == 'true',
                        reply_content=row.get('reply_content', ''),
                        reply_time=row.get('reply_time', ''),
                        created_at=row.get('created_at', ''),
                        updated_at=row.get('updated_at', ''),
                    )
                    if s.user_id:
                        self.suppliers[s.user_id] = s
            logger.info(f"已加載 {len(self.suppliers)} 個供應商記錄")
        except Exception as e:
            logger.error(f"加載數據失敗: {e}")
    
    def save(self):
        """保存數據到CSV"""
        if not self.suppliers:
            return
        fieldnames = [
            'user_id', 'nickname', 'bio', 'followers', 'notes_count', 'ip_location',
            'supplier_type', 'wechat', 'score', 'profile_url', 'keywords',
            'message_sent', 'message_content', 'message_time',
            'reply_received', 'reply_content', 'reply_time',
            'created_at', 'updated_at'
        ]
        try:
            with open(self.filepath, 'w', encoding='utf-8', newline='') as f:
                writer = csv.DictWriter(f, fieldnames=fieldnames)
                writer.writeheader()
                for s in self.suppliers.values():
                    writer.writerow(asdict(s))
            logger.info(f"已保存 {len(self.suppliers)} 個供應商記錄到 {self.filepath}")
        except Exception as e:
            logger.error(f"保存數據失敗: {e}")
    
    def add(self, supplier: Supplier) -> bool:
        """添加新供應商 (如果不存在)"""
        if supplier.user_id in self.suppliers:
            return False
        supplier.created_at = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        supplier.updated_at = supplier.created_at
        self.suppliers[supplier.user_id] = supplier
        self.save()
        return True
    
    def update(self, user_id: str, **kwargs):
        """更新供應商信息"""
        if user_id not in self.suppliers:
            return
        s = self.suppliers[user_id]
        for k, v in kwargs.items():
            if hasattr(s, k):
                setattr(s, k, v)
        s.updated_at = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        self.save()
    
    def get_unsent(self) -> List[Supplier]:
        """獲取未發送私訊的供應商"""
        return [s for s in self.suppliers.values() if not s.message_sent]
    
    def get_sent_no_reply(self) -> List[Supplier]:
        """獲取已發送但未收到回覆的供應商"""
        return [s for s in self.suppliers.values() if s.message_sent and not s.reply_received]
    
    def get_all(self) -> List[Supplier]:
        """獲取所有供應商"""
        return list(self.suppliers.values())

# ========== 工具函數 ==========
SUPPLIER_KW = ["代理", "招商", "一件代發", "批發", "工廠直供", "源頭", "廠家", "合作", "拿貨", "檔口", "一手貨源"]
FACTORY_KW = ["工廠", "廠家", "生產", "源頭"]
BRAND_KW = ["品牌", "原創", "設計師"]
EXCLUDE_KW = ["求推薦", "哪裡買", "測評", "開箱"]

WX_PATTERNS = [
    r'(?:微信|wx|vx|v信|威信|薇|徽|V)[：:\s]*([a-zA-Z0-9_\-]{5,20})',
    r'[vV][：:\s]*([a-zA-Z0-9_]{6,20})',
]

def extract_wechat(text: str) -> str:
    """提取微信號"""
    for pat in WX_PATTERNS:
        for m in re.finditer(pat, text, re.IGNORECASE):
            wxid = m.group(1).strip()
            if 5 <= len(wxid) <= 20 and any(c.isalpha() for c in wxid):
                if not any(x in wxid.lower() for x in ['xiaohongshu', 'weixin', 'wechat']):
                    return wxid
    return ""

def classify_supplier(bio: str, notes_text: str = "") -> str:
    """分類供應商類型"""
    text = (bio + " " + notes_text).lower()
    factory = sum(text.count(k) for k in FACTORY_KW)
    brand = sum(text.count(k) for k in BRAND_KW)
    if factory >= 2 and factory > brand:
        return "工廠/廠家"
    if brand >= 2:
        return "品牌方"
    if factory > 0 or brand > 0:
        return "經銷商/代理"
    return "未知"

def calc_score(supplier: Supplier) -> float:
    """計算供應商品質分數"""
    score = 0
    # 聯絡方式 (30分)
    if supplier.wechat:
        score += 30
    # 類型 (20分)
    type_scores = {"工廠/廠家": 20, "品牌方": 15, "經銷商/代理": 10}
    score += type_scores.get(supplier.supplier_type, 0)
    # 活躍度 (20分)
    if supplier.notes_count >= 50:
        score += 20
    elif supplier.notes_count >= 20:
        score += 15
    elif supplier.notes_count >= 10:
        score += 10
    elif supplier.notes_count >= 5:
        score += 5
    # 粉絲數 (15分)
    if supplier.followers >= 10000:
        score += 15
    elif supplier.followers >= 1000:
        score += 10
    elif supplier.followers >= 100:
        score += 5
    # 簡介完整度 (15分)
    if supplier.bio:
        if len(supplier.bio) >= 50:
            score += 15
        elif len(supplier.bio) >= 20:
            score += 10
        else:
            score += 5
    return round(score, 1)

def get_random_message() -> str:
    """隨機獲取一條暗語訊息"""
    return random.choice(SECRET_MESSAGES)

# ========== 爬蟲 ==========
async def delay(a=1, b=3):
    await asyncio.sleep(random.uniform(a, b))

class Scraper:
    def __init__(self, cookie: str, data_manager: DataManager):
        self.cookie = cookie
        self.dm = data_manager
        self.browser = None
        self.page = None
        self.seen = set()  # 已處理的用戶ID
        self.my_uid = None  # 當前登錄用戶ID
    
    async def init(self):
        from playwright.async_api import async_playwright
        
        logger.info("啟動瀏覽器...")
        self.pw = await async_playwright().start()
        
        # 使用持久化用戶數據目錄，保存登入狀態
        user_data_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'browser_data')
        os.makedirs(user_data_dir, exist_ok=True)
        
        # 使用 launch_persistent_context 保持登入狀態
        self.context = await self.pw.chromium.launch_persistent_context(
            user_data_dir,
            headless=HEADLESS,
            viewport={'width': 1920, 'height': 1080},
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/120.0.0.0 Safari/537.36",
            locale='zh-CN'
        )
        
        # 如果有 Cookie，也設置一下（作為備用）
        if self.cookie:
            cookies = []
            for item in self.cookie.split(';'):
                if '=' in item:
                    k, v = item.strip().split('=', 1)
                    cookies.append({'name': k, 'value': v, 'domain': '.xiaohongshu.com', 'path': '/'})
            await self.context.add_cookies(cookies)
        
        self.page = self.context.pages[0] if self.context.pages else await self.context.new_page()
        # 反偵測
        await self.page.evaluate("Object.defineProperty(navigator,'webdriver',{get:()=>undefined})")
        self.browser = None  # 使用 context 而非 browser
        logger.info("瀏覽器就緒")
    
    async def close(self):
        if self.context:
            await self.context.close()
        if hasattr(self, 'pw') and self.pw:
            await self.pw.stop()
    
    async def get_my_uid(self):
        """獲取當前登錄用戶ID"""
        try:
            await self.page.goto("https://www.xiaohongshu.com", wait_until='domcontentloaded', timeout=60000)
            await delay(2, 4)
            uid = await self.page.evaluate('''() => {
                try {
                    const state = window.__INITIAL_STATE__;
                    if (state && state.user && state.user.userPageData) {
                        return state.user.userPageData.id || '';
                    }
                    // 從 cookie 中獲取
                    const match = document.cookie.match(/userId=([^;]+)/);
                    return match ? match[1] : '';
                } catch(e) { return ''; }
            }''')
            if uid:
                self.my_uid = uid
                self.seen.add(uid)
                logger.info(f"當前用戶ID: {uid[:8]}...")
        except Exception as e:
            logger.warning(f"獲取用戶ID失敗: {e}")
    
    async def search(self, keyword: str, max_pages: int) -> List[Supplier]:
        """搜索並提取供應商"""
        found = []
        
        for pg in range(1, max_pages + 1):
            url = f"https://www.xiaohongshu.com/search_result?keyword={quote(keyword)}&source=web_search_result_notes"
            if pg > 1:
                url += f"&page={pg}"
            
            try:
                await self.page.goto(url, wait_until='domcontentloaded', timeout=60000)
                await delay(3, 5)
                
                # 滾動加載
                for _ in range(3):
                    await self.page.evaluate("window.scrollBy(0, 800)")
                    await delay(0.5, 1)
                
                logger.info(f"[{keyword}] 第{pg}/{max_pages}頁")
                
                # 提取筆記卡片
                cards = await self.page.evaluate('''() => {
                    const results = [];
                    const seen = new Set();
                    
                    // 方法1: 找所有筆記卡片鏈接
                    const noteLinks = document.querySelectorAll('a[href*="/explore/"]');
                    noteLinks.forEach(link => {
                        const href = link.getAttribute('href') || '';
                        const noteMatch = href.match(/\\/explore\\/([a-f0-9]+)/);
                        if (!noteMatch) return;
                        
                        // 找到卡片容器 - 嘗試多種方式
                        let card = link.closest('section');
                        if (!card) card = link.closest('[class*="note-item"]');
                        if (!card) card = link.closest('[class*="card"]');
                        if (!card) card = link.parentElement?.parentElement?.parentElement;
                        if (!card) return;
                        
                        // 找作者鏈接 - 可能在卡片內或附近
                        let authorId = '', authorName = '';
                        
                        // 嘗試在卡片內找
                        const authorLinks = card.querySelectorAll('a[href*="/user/profile/"]');
                        authorLinks.forEach(al => {
                            const ahref = al.getAttribute('href') || '';
                            const am = ahref.match(/\\/user\\/profile\\/([a-f0-9]+)/);
                            if (am && !authorId) {
                                const txt = (al.textContent || '').trim();
                                if (txt && txt !== '我' && txt.length > 1 && txt.length < 30) {
                                    authorId = am[1];
                                    authorName = txt;
                                }
                            }
                        });
                        
                        // 如果卡片內沒找到，嘗試從頁面狀態獲取
                        if (!authorId) {
                            try {
                                const state = window.__INITIAL_STATE__;
                                if (state) {
                                    const noteId = noteMatch[1];
                                    // 嘗試不同的數據路徑
                                    const searchNotes = state.search?.notes?.items || 
                                                       state.searchResult?.items ||
                                                       state.feed?.items || [];
                                    for (const item of searchNotes) {
                                        if (item.id === noteId || item.noteCard?.noteId === noteId) {
                                            const user = item.noteCard?.user || item.user || {};
                                            if (user.userId) {
                                                authorId = user.userId;
                                                authorName = user.nickname || '';
                                                break;
                                            }
                                        }
                                    }
                                }
                            } catch(e) {}
                        }
                        
                        if (authorId && !seen.has(authorId)) {
                            seen.add(authorId);
                            results.push({
                                noteId: noteMatch[1],
                                authorId: authorId,
                                authorName: authorName,
                                title: (card.textContent || '').slice(0, 100).replace(/\\s+/g, ' ')
                            });
                        }
                    });
                    
                    // 方法2: 如果方法1沒找到作者，嘗試從 __INITIAL_STATE__ 直接獲取
                    if (results.filter(r => r.authorId).length === 0) {
                        try {
                            const state = window.__INITIAL_STATE__;
                            if (state) {
                                const items = state.search?.notes?.items || 
                                             state.searchResult?.items ||
                                             state.feed?.items || [];
                                items.forEach(item => {
                                    const nc = item.noteCard || item;
                                    const user = nc.user || {};
                                    if (user.userId && !seen.has(user.userId)) {
                                        seen.add(user.userId);
                                        results.push({
                                            noteId: item.id || nc.noteId || '',
                                            authorId: user.userId,
                                            authorName: user.nickname || '',
                                            title: nc.displayTitle || nc.title || ''
                                        });
                                    }
                                });
                            }
                        } catch(e) {}
                    }
                    
                    return results;
                }''')
                
                logger.info(f"  找到 {len(cards)} 個筆記卡片")
                
                # 調試：顯示找到的作者ID
                author_ids = [c.get('authorId', '') for c in cards if c.get('authorId')]
                logger.info(f"  其中 {len(author_ids)} 個有作者ID")
                if author_ids[:3]:
                    logger.debug(f"  前3個作者ID: {author_ids[:3]}")
                
                # 處理每個卡片
                processed = 0
                for card in cards:
                    uid = card.get('authorId', '')
                    name = card.get('authorName', '')
                    
                    if not uid:
                        continue
                    if uid in self.seen:
                        continue
                    if self.my_uid and uid == self.my_uid:
                        continue
                    
                    self.seen.add(uid)
                    processed += 1
                    
                    # 獲取作者詳情
                    logger.info(f"  → 訪問作者: {name} ({uid[:8]}...)")
                    info = await self.get_author_info(uid, name)
                    if not info:
                        logger.warning(f"    ✗ 無法獲取作者信息")
                        continue
                    
                    logger.info(f"    獲取到: {info.get('nickname', '?')} 粉絲:{info.get('followers', 0)} 筆記:{info.get('notes', 0)}")
                    
                    # 創建供應商
                    supplier = Supplier(
                        user_id=uid,
                        nickname=info.get('nickname', name),
                        bio=info.get('bio', ''),
                        followers=info.get('followers', 0),
                        notes_count=info.get('notes', 0),
                        ip_location=info.get('ip', ''),
                        profile_url=f"https://www.xiaohongshu.com/user/profile/{uid}",
                        keywords=keyword,
                    )
                    
                    # 提取微信
                    supplier.wechat = extract_wechat(supplier.bio)
                    
                    # 分類
                    supplier.supplier_type = classify_supplier(supplier.bio)
                    
                    # 計算分數
                    supplier.score = calc_score(supplier)
                    
                    # 添加到數據管理器
                    added = self.dm.add(supplier)
                    if added:
                        found.append(supplier)
                        logger.info(f"    ✓ 已添加: {supplier.nickname} ({supplier.score}分) 微信:{supplier.wechat or '無'}")
                    else:
                        logger.info(f"    - 已存在: {supplier.nickname}")
                    
                    await delay(1, 2)
                    
                    if len(self.dm.suppliers) >= MAX_SUPPLIERS:
                        return found
                
                logger.info(f"  本頁處理了 {processed} 個新作者，當前共 {len(self.dm.suppliers)} 個供應商")
                
            except Exception as e:
                logger.error(f"搜索出錯: {e}")
            
            await delay(2, 4)
        
        return found
    
    async def get_author_info(self, uid: str, fallback_name: str = "") -> Optional[dict]:
        """獲取作者詳細信息"""
        try:
            url = f"https://www.xiaohongshu.com/user/profile/{uid}"
            resp = await self.page.goto(url, wait_until='domcontentloaded', timeout=60000)
            if resp and resp.status >= 400:
                logger.warning(f"    頁面返回錯誤: {resp.status}")
                return None
            await delay(2, 4)
            
            # 使用簡化的 JavaScript，避免正則表達式問題
            info = await self.page.evaluate("""() => {
                const result = {nickname: '', bio: '', followers: 0, notes: 0, ip: ''};
                
                // 從 __INITIAL_STATE__ 獲取
                try {
                    const state = window.__INITIAL_STATE__;
                    if (state && state.user && state.user.userPageData) {
                        const data = state.user.userPageData;
                        const basic = data.basicInfo || {};
                        result.nickname = basic.nickname || '';
                        result.bio = basic.desc || '';
                        result.ip = basic.ipLocation || '';
                        
                        const inters = data.interactions || [];
                        inters.forEach(i => {
                            if (i.type === 'fans') result.followers = parseInt(i.count) || 0;
                            if (i.type === 'notes') result.notes = parseInt(i.count) || 0;
                        });
                    }
                } catch(e) {}
                
                // 備用: 從 DOM 獲取
                if (!result.nickname) {
                    const nameEl = document.querySelector('.user-name, [class*="nickname"], h1');
                    result.nickname = nameEl ? nameEl.textContent.trim() : '';
                }
                if (!result.bio) {
                    const bioEl = document.querySelector('.user-desc, [class*="desc"], [class*="bio"]');
                    result.bio = bioEl ? bioEl.textContent.trim() : '';
                }
                
                // 從頁面文本解析粉絲數和筆記數
                const pageText = document.body.innerText || '';
                
                if (!result.followers) {
                    // 嘗試找「X 粉絲」或「X粉絲」
                    const textParts = pageText.split(/\\s+/);
                    for (let i = 0; i < textParts.length; i++) {
                        if (textParts[i].includes('粉絲') || textParts[i].includes('粉丝')) {
                            // 檢查前一個詞是否是數字
                            if (i > 0) {
                                let numStr = textParts[i-1].replace(/[,，]/g, '');
                                if (numStr.includes('万') || numStr.includes('萬')) {
                                    result.followers = Math.round(parseFloat(numStr) * 10000);
                                } else if (numStr.toLowerCase().includes('k')) {
                                    result.followers = Math.round(parseFloat(numStr) * 1000);
                                } else {
                                    result.followers = parseInt(numStr) || 0;
                                }
                            }
                            break;
                        }
                    }
                }
                
                if (!result.notes) {
                    const textParts = pageText.split(/\\s+/);
                    for (let i = 0; i < textParts.length; i++) {
                        if (textParts[i].includes('筆記') || textParts[i].includes('笔记')) {
                            if (i > 0) {
                                result.notes = parseInt(textParts[i-1].replace(/[,，]/g, '')) || 0;
                            }
                            break;
                        }
                    }
                }
                
                if (!result.ip) {
                    const idx = pageText.indexOf('IP');
                    if (idx > -1) {
                        // 找 IP屬地: 或 IP属地: 後面的文字
                        const after = pageText.slice(idx, idx + 30);
                        const parts = after.split(/[：:]/);
                        if (parts.length > 1) {
                            result.ip = parts[1].trim().split(/\\s/)[0] || '';
                        }
                    }
                }
                
                return result;
            }""")
            
            if not info.get('nickname'):
                info['nickname'] = fallback_name
            
            # 調試輸出
            logger.debug(f"    提取結果: nickname={info.get('nickname')}, bio長度={len(info.get('bio', ''))}")
            
            return info if info.get('nickname') else None
            
        except Exception as e:
            logger.warning(f"    獲取作者信息異常: {e}")
            return None
    
    async def send_message(self, supplier: Supplier, message: str) -> bool:
        """發送私訊給供應商"""
        try:
            logger.info(f"準備發送私訊給: {supplier.nickname}")
            
            # 訪問用戶主頁
            await self.page.goto(supplier.profile_url, wait_until='domcontentloaded', timeout=60000)
            await delay(3, 5)
            
            # 嘗試多種方式找到並點擊私信按鈕
            clicked = await self.page.evaluate("""() => {
                // 方法1: 找所有元素，匹配私信相關文字
                const allElements = document.querySelectorAll('button, div, span, a, li');
                
                for (const el of allElements) {
                    const text = (el.textContent || '').trim();
                    // 匹配各種私信按鈕名稱
                    if (text === '私信' || text === '发私信' || text === '發私信' || 
                        text === '发消息' || text === '發消息' || text === '聊天') {
                        const rect = el.getBoundingClientRect();
                        if (rect.width > 0 && rect.height > 0) {
                            el.click();
                            return {success: true, text: text, method: 'text'};
                        }
                    }
                }
                
                // 方法2: 找用戶操作區域 (關注按鈕旁邊通常有私信)
                const followBtn = Array.from(document.querySelectorAll('button, div, span')).find(
                    el => el.textContent.trim() === '关注' || el.textContent.trim() === '關注'
                );
                if (followBtn) {
                    // 找附近的私信按鈕
                    const parent = followBtn.parentElement;
                    if (parent) {
                        const siblings = parent.querySelectorAll('button, div, span, a');
                        for (const sib of siblings) {
                            const t = (sib.textContent || '').trim();
                            if (t.includes('私信') || t.includes('消息') || t.includes('聊')) {
                                sib.click();
                                return {success: true, text: t, method: 'sibling'};
                            }
                        }
                    }
                }
                
                // 方法3: 找 SVG 圖標按鈕 (有些按鈕只有圖標沒有文字)
                const svgBtns = document.querySelectorAll('svg, [class*="icon"]');
                svgBtns.forEach(svg => {
                    const parent = svg.closest('button, div, a, span');
                    if (parent) {
                        const cls = (parent.className || '').toLowerCase();
                        if (cls.includes('message') || cls.includes('chat') || cls.includes('dm')) {
                            parent.click();
                            return {success: true, method: 'icon', class: cls};
                        }
                    }
                });
                
                // 方法4: 嘗試點擊「更多」按鈕展開菜單
                const moreBtn = Array.from(document.querySelectorAll('button, div, span')).find(
                    el => {
                        const t = el.textContent.trim();
                        return t === '更多' || t === '...' || t === '⋯';
                    }
                );
                if (moreBtn) {
                    moreBtn.click();
                    return {success: false, needWait: true, message: '點擊了更多按鈕'};
                }
                
                // 調試: 返回頁面上所有可點擊元素的文字
                const btnTexts = [];
                document.querySelectorAll('button, [role="button"], [class*="btn"], a').forEach(b => {
                    const t = (b.textContent || '').trim();
                    if (t && t.length < 30 && !btnTexts.includes(t)) btnTexts.push(t);
                });
                
                return {success: false, buttons: btnTexts.slice(0, 20)};
            }""")
            
            if not clicked.get('success'):
                # 如果點擊了「更多」按鈕，等一下再試
                if clicked.get('needWait'):
                    await delay(1, 2)
                    # 再次嘗試找私信按鈕
                    clicked2 = await self.page.evaluate("""() => {
                        const allElements = document.querySelectorAll('button, div, span, a, li');
                        for (const el of allElements) {
                            const text = (el.textContent || '').trim();
                            if (text === '私信' || text === '发私信' || text === '發私信' ||
                                text === '发消息' || text === '發消息') {
                                el.click();
                                return {success: true, text: text};
                            }
                        }
                        return {success: false};
                    }""")
                    if clicked2.get('success'):
                        clicked = clicked2
                    else:
                        logger.warning(f"找不到私信按鈕: {supplier.nickname}")
                        return False
                else:
                    logger.warning(f"找不到私信按鈕: {supplier.nickname}")
                    btns = clicked.get('buttons', [])
                    if btns:
                        logger.info(f"  頁面按鈕: {btns[:10]}")
                    # 保存截圖幫助調試
                    try:
                        await self.page.screenshot(path=f"debug_profile_{supplier.user_id[:8]}.png")
                        logger.info(f"  已保存截圖: debug_profile_{supplier.user_id[:8]}.png")
                    except:
                        pass
                    return False
            
            logger.info(f"  ✓ 點擊私信按鈕")
            await delay(2, 4)
            
            # 等待聊天窗口出現
            await self.page.wait_for_timeout(2000)
            
            # 嘗試多種輸入框選擇器
            input_found = await self.page.evaluate("""() => {
                const selectors = [
                    'textarea',
                    '[contenteditable="true"]',
                    'input[type="text"]',
                    '[class*="editor"]',
                    '[class*="input"]',
                    '[class*="textarea"]'
                ];
                
                for (const sel of selectors) {
                    const inputs = document.querySelectorAll(sel);
                    for (const inp of inputs) {
                        const rect = inp.getBoundingClientRect();
                        if (rect.width > 50 && rect.height > 15) {
                            inp.focus();
                            inp.click();
                            return {found: true, selector: sel};
                        }
                    }
                }
                return {found: false};
            }""")
            
            if not input_found.get('found'):
                logger.warning(f"找不到輸入框: {supplier.nickname}")
                return False
            
            logger.info(f"  ✓ 找到輸入框")
            await delay(0.5, 1)
            
            # 使用鍵盤輸入訊息 (更可靠)
            await self.page.keyboard.type(message, delay=30)
            await delay(1, 2)
            
            # 發送訊息
            send_result = await self.page.evaluate("""() => {
                const btns = document.querySelectorAll('button, div, span');
                for (const btn of btns) {
                    const text = (btn.textContent || '').trim();
                    if (text === '发送' || text === '發送') {
                        const rect = btn.getBoundingClientRect();
                        if (rect.width > 0) {
                            btn.click();
                            return {clicked: true};
                        }
                    }
                }
                return {clicked: false};
            }""")
            
            if not send_result.get('clicked'):
                # 嘗試按 Enter 發送
                await self.page.keyboard.press('Enter')
            
            await delay(2, 3)
            
            # 更新數據
            self.dm.update(
                supplier.user_id,
                message_sent=True,
                message_content=message,
                message_time=datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            )
            
            logger.info(f"✓ 已發送私訊給: {supplier.nickname}")
            return True
            
        except Exception as e:
            logger.error(f"發送私訊失敗 {supplier.nickname}: {e}")
            return False
    
    async def check_replies(self) -> List[dict]:
        """檢查私訊回覆"""
        replies = []
        
        try:
            logger.info("檢查私訊回覆...")
            
            # 訪問消息頁面
            await self.page.goto("https://www.xiaohongshu.com/notification", wait_until='domcontentloaded', timeout=60000)
            await delay(3, 5)
            
            # 點擊私信標籤
            dm_tab = await self.page.query_selector('span:has-text("私信"), [class*="message"], [class*="chat"]')
            if dm_tab:
                await dm_tab.click()
                await delay(2, 4)
            
            # 提取對話列表
            conversations = await self.page.evaluate('''() => {
                const results = [];
                // 查找對話項目
                const items = document.querySelectorAll('[class*="conversation"], [class*="chat-item"], [class*="message-item"]');
                items.forEach(item => {
                    const nameEl = item.querySelector('[class*="name"], [class*="nickname"]');
                    const msgEl = item.querySelector('[class*="content"], [class*="message"], [class*="text"]');
                    const linkEl = item.querySelector('a[href*="/user/profile/"]');
                    
                    let uid = '';
                    if (linkEl) {
                        const match = linkEl.href.match(/\\/user\\/profile\\/([a-f0-9]+)/);
                        if (match) uid = match[1];
                    }
                    
                    if (nameEl) {
                        results.push({
                            nickname: nameEl.textContent.trim(),
                            lastMessage: msgEl ? msgEl.textContent.trim() : '',
                            userId: uid
                        });
                    }
                });
                return results;
            }''')
            
            logger.info(f"找到 {len(conversations)} 個對話")
            
            # 獲取已發送私訊的用戶
            sent_suppliers = {s.user_id: s for s in self.dm.get_sent_no_reply()}
            
            for conv in conversations:
                uid = conv.get('userId', '')
                nickname = conv.get('nickname', '')
                last_msg = conv.get('lastMessage', '')
                
                # 檢查是否是我們發過私訊的用戶
                supplier = None
                if uid and uid in sent_suppliers:
                    supplier = sent_suppliers[uid]
                else:
                    # 嘗試通過暱稱匹配
                    for s in sent_suppliers.values():
                        if s.nickname == nickname:
                            supplier = s
                            break
                
                if supplier and last_msg:
                    # 檢查是否是新回覆 (不是我們發的訊息)
                    is_our_msg = any(m in last_msg for m in SECRET_MESSAGES[:5])  # 檢查前幾條模板
                    if not is_our_msg:
                        # 更新回覆
                        self.dm.update(
                            supplier.user_id,
                            reply_received=True,
                            reply_content=last_msg[:500],
                            reply_time=datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                        )
                        replies.append({
                            'nickname': supplier.nickname,
                            'user_id': supplier.user_id,
                            'reply': last_msg
                        })
                        logger.info(f"✓ 收到回覆: {supplier.nickname} - {last_msg[:50]}...")
            
        except Exception as e:
            logger.error(f"檢查回覆失敗: {e}")
        
        return replies
    
    async def run_search(self, keywords: list[str], max_pages: int):
        """運行搜索任務"""
        await self.init()
        try:
            await self.get_my_uid()
            
            # 加載已有的用戶ID到 seen 集合
            for uid in self.dm.suppliers.keys():
                self.seen.add(uid)
            
            for kw in keywords:
                logger.info(f"搜索: {kw}")
                await self.search(kw, max_pages)
                
                if len(self.dm.suppliers) >= MAX_SUPPLIERS:
                    logger.info(f"已達到最大供應商數量 {MAX_SUPPLIERS}")
                    break
                
                await delay(3, 5)
        finally:
            await self.close()
    
    async def run_messaging(self, max_messages: int = MAX_MESSAGES_PER_RUN):
        """運行私訊任務 (低頻率)"""
        await self.init()
        try:
            await self.get_my_uid()
            
            # 獲取未發送私訊的供應商
            unsent = self.dm.get_unsent()
            if not unsent:
                logger.info("沒有需要發送私訊的供應商")
                return
            
            # 隨機打亂順序
            random.shuffle(unsent)
            
            sent_count = 0
            for supplier in unsent:
                if sent_count >= max_messages:
                    logger.info(f"已達到本次最大私訊數量 {max_messages}")
                    break
                
                # 隨機選擇暗語訊息
                message = get_random_message()
                
                logger.info(f"\n{'='*40}")
                logger.info(f"[{sent_count+1}/{max_messages}] 準備私訊: {supplier.nickname}")
                logger.info(f"暗語: {message}")
                
                success = await self.send_message(supplier, message)
                
                if success:
                    sent_count += 1
                    
                    if sent_count < max_messages and sent_count < len(unsent):
                        # 低頻率延遲 (5-10分鐘)
                        wait_time = random.randint(MESSAGE_DELAY_MIN, MESSAGE_DELAY_MAX)
                        logger.info(f"等待 {wait_time//60} 分 {wait_time%60} 秒後發送下一條...")
                        await asyncio.sleep(wait_time)
                
                await delay(2, 4)
            
            logger.info(f"\n本次共發送 {sent_count} 條私訊")
            
        finally:
            await self.close()
    
    async def run_check_replies(self):
        """運行回覆檢查任務"""
        await self.init()
        try:
            replies = await self.check_replies()
            if replies:
                logger.info(f"\n收到 {len(replies)} 條新回覆:")
                for r in replies:
                    logger.info(f"  - {r['nickname']}: {r['reply'][:80]}...")
            else:
                logger.info("暫無新回覆")
        finally:
            await self.close()

# ========== 導出 ==========
def export_excel(dm: DataManager, filename: str = None):
    """導出到Excel"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        logger.error("請安裝 openpyxl: pip install openpyxl")
        return None
    
    suppliers = sorted(dm.get_all(), key=lambda s: s.score, reverse=True)
    if not suppliers:
        logger.warning("沒有數據可導出")
        return None
    
    if not filename:
        filename = f"童裝供應商_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    wb = Workbook()
    ws = wb.active
    ws.title = "供應商列表"
    
    headers = [
        "排名", "暱稱", "品質評分", "供應商類型", "微信", 
        "粉絲數", "筆記數", "IP屬地", "主頁連結", "簡介",
        "私訊狀態", "發送時間", "發送內容",
        "回覆狀態", "回覆時間", "回覆內容"
    ]
    widths = [6, 16, 10, 12, 20, 10, 8, 10, 45, 35, 10, 18, 40, 10, 18, 50]
    
    # 表頭
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(1, col, h)
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[get_column_letter(col)].width = w
    
    ws.freeze_panes = 'A2'
    
    # 數據
    for i, s in enumerate(suppliers, 1):
        row_data = [
            i,
            s.nickname,
            s.score,
            s.supplier_type,
            s.wechat or "未找到",
            s.followers,
            s.notes_count,
            s.ip_location,
            s.profile_url,
            (s.bio[:80] + "...") if len(s.bio) > 80 else s.bio,
            "已發送" if s.message_sent else "未發送",
            s.message_time,
            (s.message_content[:40] + "...") if len(s.message_content) > 40 else s.message_content,
            "已回覆" if s.reply_received else ("待回覆" if s.message_sent else "-"),
            s.reply_time,
            (s.reply_content[:50] + "...") if len(s.reply_content) > 50 else s.reply_content,
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(i + 1, col, val)
            if "連結" in headers[col-1] and val:
                cell.hyperlink = val
                cell.font = Font(color="0563C1", underline="single")
            # 高亮已回覆的行
            if s.reply_received:
                cell.fill = PatternFill("solid", fgColor="C6EFCE")
    
    os.makedirs("output", exist_ok=True)
    path = f"output/{filename}"
    wb.save(path)
    logger.info(f"已導出: {path}")
    return path

# ========== 主程序 ==========
def main():
    print("""
╔═══════════════════════════════════════════════════════╗
║     小紅書童裝供應商爬蟲 v2.0 (暗語私訊版)            ║
╚═══════════════════════════════════════════════════════╝
    """)
    
    global COOKIE
    
    # 初始化數據管理器
    dm = DataManager(DATA_FILE)
    
    # 獲取 Cookie
    # Cookie 現在是可選的，因為使用持久化登入
    if not COOKIE:
        print("💡 提示：首次使用請選擇「0. 手動登入」")
        print("   如果已登入過，可直接選擇其他操作")
        print()
        print("   (可選) 輸入 Cookie 或直接按 Enter 跳過：")
        COOKIE = input("Cookie: ").strip()
        if not COOKIE:
            print("   → 跳過 Cookie，將使用已保存的登入狀態")
    
    while True:
        print("\n" + "="*50)
        print("請選擇操作:")
        print("  0. 🔐 手動登入小紅書 (首次使用必選)")
        print("  1. 🔍 搜索供應商")
        print("  2. 💬 發送私訊 (暗語)")
        print("  3. 📥 檢查回覆")
        print("  4. 📊 導出Excel")
        print("  5. 📋 查看統計")
        print("  6. 🚪 退出")
        print("="*50)
        
        choice = input("\n請輸入選項 (0-6): ").strip()
        
        if choice == '0':
            # 手動登入
            print("\n🔐 手動登入模式")
            print("將打開瀏覽器，請手動掃碼登入小紅書")
            print("登入成功後，關閉瀏覽器窗口即可")
            input("\n按 Enter 打開瀏覽器...")
            
            async def manual_login():
                from playwright.async_api import async_playwright
                
                user_data_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'browser_data')
                os.makedirs(user_data_dir, exist_ok=True)
                
                pw = await async_playwright().start()
                context = await pw.chromium.launch_persistent_context(
                    user_data_dir,
                    headless=False,  # 必須顯示窗口
                    viewport={'width': 1280, 'height': 800},
                    user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/120.0.0.0 Safari/537.36",
                    locale='zh-CN'
                )
                
                page = context.pages[0] if context.pages else await context.new_page()
                await page.goto("https://www.xiaohongshu.com")
                
                print("\n⏳ 請在瀏覽器中登入小紅書...")
                print("   登入成功後，請關閉瀏覽器窗口")
                
                # 等待用戶關閉瀏覽器
                try:
                    await page.wait_for_event('close', timeout=300000)  # 5分鐘超時
                except:
                    pass
                
                await context.close()
                await pw.stop()
                print("\n✅ 登入狀態已保存！")
            
            try:
                asyncio.run(manual_login())
            except Exception as e:
                print(f"\n錯誤: {e}")
        
        elif choice == '1':
            # 搜索供應商
            print(f"\n🔍 關鍵字: {KEYWORDS}")
            print(f"⏱️ 每個關鍵字 {MAX_PAGES} 頁")
            input("\n按 Enter 開始搜索...")
            
            scraper = Scraper(COOKIE, dm)
            try:
                asyncio.run(scraper.run_search(KEYWORDS, MAX_PAGES))
            except KeyboardInterrupt:
                print("\n用戶中斷")
            except Exception as e:
                print(f"\n錯誤: {e}")
            
            print(f"\n✅ 當前共 {len(dm.suppliers)} 個供應商")
        
        elif choice == '2':
            # 發送私訊
            unsent = dm.get_unsent()
            if not unsent:
                print("\n沒有需要發送私訊的供應商")
                print("請先執行搜索操作 (選項1)")
                continue
            
            print(f"\n📮 待發送私訊: {len(unsent)} 個供應商")
            print(f"⏱️ 發送間隔: {MESSAGE_DELAY_MIN//60}-{MESSAGE_DELAY_MAX//60} 分鐘")
            print(f"📝 本次最多發送: {MAX_MESSAGES_PER_RUN} 條")
            print("\n⚠️ 暗語訊息將隨機選擇，降低被偵測風險")
            print("\n暗語示例:")
            for i, msg in enumerate(random.sample(SECRET_MESSAGES, min(3, len(SECRET_MESSAGES))), 1):
                print(f"  {i}. {msg}")
            
            input("\n按 Enter 開始發送...")
            
            scraper = Scraper(COOKIE, dm)
            try:
                asyncio.run(scraper.run_messaging(MAX_MESSAGES_PER_RUN))
            except KeyboardInterrupt:
                print("\n用戶中斷")
            except Exception as e:
                print(f"\n錯誤: {e}")
        
        elif choice == '3':
            # 檢查回覆
            sent = dm.get_sent_no_reply()
            if not sent:
                print("\n沒有待檢查回覆的供應商")
                continue
            
            print(f"\n📥 檢查 {len(sent)} 個已發送私訊的供應商的回覆")
            input("\n按 Enter 開始檢查...")
            
            scraper = Scraper(COOKIE, dm)
            try:
                asyncio.run(scraper.run_check_replies())
            except KeyboardInterrupt:
                print("\n用戶中斷")
            except Exception as e:
                print(f"\n錯誤: {e}")
        
        elif choice == '4':
            # 導出Excel
            if not dm.suppliers:
                print("\n沒有數據可導出")
                continue
            
            path = export_excel(dm)
            if path:
                print(f"\n📊 已導出: {path}")
                print(f"📂 打開: open output/")
        
        elif choice == '5':
            # 查看統計
            all_suppliers = dm.get_all()
            sent = [s for s in all_suppliers if s.message_sent]
            replied = [s for s in all_suppliers if s.reply_received]
            with_wechat = [s for s in all_suppliers if s.wechat]
            
            print("\n📋 數據統計:")
            print(f"   總供應商數: {len(all_suppliers)}")
            print(f"   有微信號: {len(with_wechat)}")
            print(f"   已發私訊: {len(sent)}")
            print(f"   已收回覆: {len(replied)}")
            
            if replied:
                print("\n📬 已回覆供應商:")
                for s in replied[:10]:
                    print(f"   - {s.nickname}: {s.reply_content[:40]}...")
        
        elif choice == '6':
            print("\n👋 再見!")
            break
        
        else:
            print("\n⚠️ 無效選項，請輸入 0-6")

if __name__ == "__main__":
    main()
